import os
import hashlib
import aiofiles
from typing import List, Dict, Optional, Tuple
from pathlib import Path
import asyncio
from datetime import datetime

# Document processing libraries
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from bs4 import BeautifulSoup
import requests
from readability import Document as ReadabilityDocument
import csv
import json as json_module

from app.core.config import settings
from app.core.exceptions import FileProcessingError
from app.core.security import create_secure_filename


class DocumentProcessor:
    """Document processing service for handling various file formats."""

    def __init__(self):
        self.supported_extensions = {
            '.pdf', '.docx', '.doc', '.txt', '.md',
            '.html', '.htm', '.csv', '.json', '.xml'
        }

    async def process_file(
        self,
        file_path: str,
        original_filename: str,
        user_id: int
    ) -> Dict[str, any]:
        """Process uploaded file and extract content."""
        try:
            # Validate file
            file_extension = Path(original_filename).suffix.lower()
            if file_extension not in self.supported_extensions:
                raise FileProcessingError(f"Unsupported file type: {file_extension}")

            # Calculate file hash
            file_hash = await self._calculate_file_hash(file_path)

            # Get file size
            file_size = os.path.getsize(file_path)

            # Determine MIME type
            mime_type = self._get_mime_type(file_extension)

            # Extract content based on file type
            content, metadata = await self._extract_content(file_path, file_extension)

            # Generate chunks for processing
            chunks = self._chunk_text(content)

            # Extract additional metadata
            word_count = len(content.split())
            metadata.update({
                'word_count': word_count,
                'processed_at': datetime.utcnow().isoformat(),
                'chunk_count': len(chunks)
            })

            return {
                'content': content,
                'chunks': chunks,
                'metadata': metadata,
                'file_hash': file_hash,
                'file_size': file_size,
                'mime_type': mime_type,
                'word_count': word_count
            }

        except Exception as e:
            raise FileProcessingError(f"Failed to process file: {str(e)}")

    async def _calculate_file_hash(self, file_path: str) -> str:
        """Calculate SHA-256 hash of file."""
        sha256_hash = hashlib.sha256()
        async with aiofiles.open(file_path, 'rb') as f:
            async for chunk in f:
                sha256_hash.update(chunk)
        return sha256_hash.hexdigest()

    def _get_mime_type(self, file_extension: str) -> str:
        """Get MIME type from file extension."""
        mime_types = {
            '.pdf': 'application/pdf',
            '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            '.doc': 'application/msword',
            '.txt': 'text/plain',
            '.md': 'text/markdown',
            '.html': 'text/html',
            '.htm': 'text/html',
            '.csv': 'text/csv',
            '.json': 'application/json',
            '.xml': 'application/xml'
        }
        return mime_types.get(file_extension, 'application/octet-stream')

    async def _extract_content(self, file_path: str, file_extension: str) -> Tuple[str, Dict]:
        """Extract content from file based on type."""
        processors = {
            '.pdf': self._extract_pdf_content,
            '.docx': self._extract_docx_content,
            '.doc': self._extract_doc_content,
            '.txt': self._extract_text_content,
            '.md': self._extract_markdown_content,
            '.html': self._extract_html_content,
            '.htm': self._extract_html_content,
            '.csv': self._extract_csv_content,
            '.json': self._extract_json_content,
            '.xml': self._extract_xml_content
        }

        processor = processors.get(file_extension)
        if not processor:
            raise FileProcessingError(f"No processor for file type: {file_extension}")

        return await processor(file_path)

    async def _extract_pdf_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from PDF file."""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PdfReader(file)
                content = ""
                metadata = {
                    'page_count': len(pdf_reader.pages),
                    'title': pdf_reader.metadata.get('/Title', '') if pdf_reader.metadata else '',
                    'author': pdf_reader.metadata.get('/Author', '') if pdf_reader.metadata else ''
                }

                for page_num, page in enumerate(pdf_reader.pages, 1):
                    try:
                        page_text = page.extract_text()
                        content += f"\n\n--- Page {page_num} ---\n{page_text}"
                    except Exception as e:
                        print(f"Error extracting page {page_num}: {e}")
                        continue

                return content.strip(), metadata

        except Exception as e:
            raise FileProcessingError(f"Failed to extract PDF content: {str(e)}")

    async def _extract_docx_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from DOCX file."""
        try:
            doc = DocxDocument(file_path)
            content = "\n".join([paragraph.text for paragraph in doc.paragraphs])

            # Extract metadata
            metadata = {
                'title': doc.core_properties.title or '',
                'author': doc.core_properties.author or '',
                'created': doc.core_properties.created.isoformat() if doc.core_properties.created else ''
            }

            return content.strip(), metadata

        except Exception as e:
            raise FileProcessingError(f"Failed to extract DOCX content: {str(e)}")

    async def _extract_doc_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from legacy DOC file."""
        # Note: DOC files require more complex processing. For now, we'll use a placeholder.
        # In production, you might want to use antiword or similar tools.
        raise FileProcessingError("Legacy DOC files not yet supported. Please convert to DOCX.")

    async def _extract_text_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from plain text file."""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                content = await file.read()

            metadata = {
                'encoding': 'utf-8',
                'line_count': content.count('\n') + 1
            }

            return content.strip(), metadata

        except UnicodeDecodeError:
            # Try with different encoding
            try:
                async with aiofiles.open(file_path, 'r', encoding='latin-1') as file:
                    content = await file.read()
                metadata = {'encoding': 'latin-1', 'line_count': content.count('\n') + 1}
                return content.strip(), metadata
            except Exception as e:
                raise FileProcessingError(f"Failed to read text file with any encoding: {str(e)}")

        except Exception as e:
            raise FileProcessingError(f"Failed to extract text content: {str(e)}")

    async def _extract_markdown_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from Markdown file."""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                content = await file.read()

            metadata = {
                'format': 'markdown',
                'line_count': content.count('\n') + 1
            }

            return content.strip(), metadata

        except Exception as e:
            raise FileProcessingError(f"Failed to extract Markdown content: {str(e)}")

    async def _extract_html_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from HTML file."""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                html_content = await file.read()

            # Use BeautifulSoup to extract text
            soup = BeautifulSoup(html_content, 'html.parser')

            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()

            content = soup.get_text(separator=' ', strip=True)

            # Extract metadata
            title = soup.title.string if soup.title else ''
            metadata = {
                'title': title.strip(),
                'format': 'html'
            }

            return content.strip(), metadata

        except Exception as e:
            raise FileProcessingError(f"Failed to extract HTML content: {str(e)}")

    async def _extract_csv_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from CSV file."""
        try:
            content_parts = []
            metadata = {'format': 'csv', 'columns': []}

            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                content = await file.read()

            # Parse CSV to get structure
            lines = content.split('\n')
            if lines:
                # Try to detect delimiter
                sample = lines[0]
                if ',' in sample:
                    delimiter = ','
                elif ';' in sample:
                    delimiter = ';'
                elif '\t' in sample:
                    delimiter = '\t'
                else:
                    delimiter = ','

                try:
                    reader = csv.reader(content.splitlines(), delimiter=delimiter)
                    headers = next(reader, [])
                    metadata['columns'] = headers
                    metadata['row_count'] = len(lines) - 1

                    # Convert CSV to readable text
                    for i, row in enumerate(reader):
                        if i < 100:  # Limit to first 100 rows for processing
                            content_parts.append(f"Row {i+1}: {' | '.join(row)}")

                except Exception:
                    # Fallback: treat as plain text
                    content_parts = lines[:100]  # Limit to first 100 lines

            return '\n'.join(content_parts), metadata

        except Exception as e:
            raise FileProcessingError(f"Failed to extract CSV content: {str(e)}")

    async def _extract_json_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from JSON file."""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                json_content = await file.read()

            try:
                data = json_module.loads(json_content)
                content = json_module.dumps(data, indent=2, ensure_ascii=False)

                metadata = {
                    'format': 'json',
                    'keys': self._extract_json_keys(data)
                }

                return content, metadata

            except json_module.JSONDecodeError as e:
                raise FileProcessingError(f"Invalid JSON format: {str(e)}")

        except Exception as e:
            raise FileProcessingError(f"Failed to extract JSON content: {str(e)}")

    async def _extract_xml_content(self, file_path: str) -> Tuple[str, Dict]:
        """Extract content from XML file."""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                xml_content = await file.read()

            soup = BeautifulSoup(xml_content, 'xml')
            content = soup.get_text(separator=' ', strip=True)

            metadata = {
                'format': 'xml',
                'root_element': soup.root.name if soup.root else ''
            }

            return content.strip(), metadata

        except Exception as e:
            raise FileProcessingError(f"Failed to extract XML content: {str(e)}")

    def _extract_json_keys(self, obj, max_depth=3, current_depth=0) -> List[str]:
        """Recursively extract keys from JSON object."""
        if current_depth >= max_depth:
            return []

        keys = []
        if isinstance(obj, dict):
            for key, value in obj.items():
                keys.append(key)
                if isinstance(value, (dict, list)):
                    keys.extend(self._extract_json_keys(value, max_depth, current_depth + 1))
        elif isinstance(obj, list):
            for item in obj:
                if isinstance(item, (dict, list)):
                    keys.extend(self._extract_json_keys(item, max_depth, current_depth + 1))

        return list(set(keys))  # Remove duplicates

    def _chunk_text(self, text: str, chunk_size: int = None, overlap: int = None) -> List[str]:
        """Split text into chunks for processing."""
        if chunk_size is None:
            chunk_size = settings.CHUNK_SIZE
        if overlap is None:
            overlap = settings.CHUNK_OVERLAP

        if len(text) <= chunk_size:
            return [text]

        chunks = []
        start = 0

        while start < len(text):
            end = start + chunk_size

            # Try to break at sentence or paragraph boundaries
            if end < len(text):
                # Look for sentence endings
                sentence_end = max(
                    text.rfind('. ', start, end),
                    text.rfind('! ', start, end),
                    text.rfind('? ', start, end)
                )

                if sentence_end > start + chunk_size // 2:  # Found good sentence break
                    end = sentence_end + 2
                else:
                    # Look for paragraph break
                    paragraph_end = text.rfind('\n\n', start, end)
                    if paragraph_end > start + chunk_size // 2:
                        end = paragraph_end + 2

            chunks.append(text[start:end].strip())
            start = max(start + 1, end - overlap)

        return [chunk for chunk in chunks if chunk]

    async def process_url(self, url: str) -> Dict[str, any]:
        """Process content from a URL."""
        try:
            # Fetch webpage content
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
            }

            response = requests.get(url, headers=headers, timeout=30)
            response.raise_for_status()

            # Use readability to extract main content
            doc = ReadabilityDocument(response.text)
            content = doc.summary()
            title = doc.title()

            # Generate chunks
            chunks = self._chunk_text(content)

            metadata = {
                'source_url': url,
                'title': title,
                'processed_at': datetime.utcnow().isoformat(),
                'word_count': len(content.split()),
                'chunk_count': len(chunks)
            }

            return {
                'content': content,
                'chunks': chunks,
                'metadata': metadata,
                'title': title
            }

        except Exception as e:
            raise FileProcessingError(f"Failed to process URL: {str(e)}")

    def get_supported_formats(self) -> List[str]:
        """Get list of supported file formats."""
        return list(self.supported_extensions)

    def is_format_supported(self, file_extension: str) -> bool:
        """Check if file format is supported."""
        return file_extension.lower() in self.supported_extensions